"""Batch feature/correlation compares driven by an Excel pair list."""

from __future__ import annotations

from pathlib import Path
from typing import Any, BinaryIO

from compare import (
    _attach_correlation_stats,
    build_feature_compare,
    build_feature_compare_ze2,
)
from features import analyze_signal_features, compare_feature_rows
from parsers.txt_device import TXT_MV_PER_COUNT, load_txt_emg
from parsers.ze2_txt import DEFAULT_SAMPLE_RATE as ZE2_DEFAULT_FS
from parsers.ze2_txt import ZE2_MV_PER_COUNT
from paths import ROOT

# Analysis recipe for ZE1 list runs. Fatigue uses 10 contractions; others use 3.
ANALYSIS_CONFIG: dict[str, Any] = {
    "contraction_method": "ze1_schmitt",
    "contraction_method_label": "ZE1 施密特觸發（Schmitt trigger）",
    "feature_method": "ttri",
    "feature_method_label": "TTRI（AEMG + 滑動窗 RMS/iEMG/MPF/MDF）",
    "expected_count_default": 3,
    "expected_count_by_purpose": {
        "裝置比對": 3,
        "刮腿毛": 3,
        "疲勞": 10,
    },
    "ze2_sample_rate_hz": float(ZE2_DEFAULT_FS),
    "ze2_mv_per_count": float(ZE2_MV_PER_COUNT),
}


def expected_count_for(purpose: str) -> int:
    mapping = ANALYSIS_CONFIG.get("expected_count_by_purpose") or {}
    if purpose in mapping:
        return int(mapping[purpose])
    return int(ANALYSIS_CONFIG.get("expected_count_default") or 3)


def index_data_files(root: Path | None = None) -> dict[str, Path]:
    base = root or (ROOT / "data")
    by_lower: dict[str, Path] = {}
    for path in base.rglob("*"):
        if path.is_file():
            by_lower.setdefault(path.name.lower(), path)
    return by_lower


def resolve_name(name: str | None, index: dict[str, Path], *, root: Path | None = None) -> Path | None:
    if not name or not str(name).strip():
        return None
    raw = str(name).strip()
    base = root or (ROOT / "data")
    exact = list(base.rglob(raw))
    if exact:
        return exact[0]
    return index.get(raw.lower())


def build_feature_compare_txt_pair(
    left_name: str,
    right_name: str,
    *,
    expected_count: int = 3,
    contraction_method: str = "ze1_schmitt",
    feature_method: str = "ttri",
) -> dict[str, Any]:
    """ZE1 TXT vs ZE1 TXT (e.g. unshaved vs shaved)."""
    left = load_txt_emg(left_name, for_plot=False)
    right = load_txt_emg(right_name, for_plot=False)
    left_feat = analyze_signal_features(
        left["times"],
        left["values"],
        sample_rate=left["sample_rate"],
        expected_count=expected_count,
        contraction_method=contraction_method,
        feature_method=feature_method,
        source="txt",
    )
    right_feat = analyze_signal_features(
        right["times"],
        right["values"],
        sample_rate=right["sample_rate"],
        expected_count=expected_count,
        contraction_method=contraction_method,
        feature_method=feature_method,
        source="txt",
    )
    pairs = compare_feature_rows(
        left_feat["features"],
        right_feat["features"],
        metrics=left_feat["metrics"],
    )
    result = {
        "mode": "features",
        "device": "ze1_vs_ze1",
        "expected_count": expected_count,
        "contraction_method": contraction_method,
        "feature_method": feature_method,
        "metrics": left_feat["metrics"],
        "delsys": {
            "filename": left["filename"],
            "signal_name": left["signal_name"],
            "sample_rate": left["sample_rate"],
            "unit": left["unit"],
            "features": left_feat["features"],
            "count": left_feat["count"],
            "role": "reference_ze1",
        },
        "txt": {
            "filename": right["filename"],
            "signal_name": right["signal_name"],
            "sample_rate": right["sample_rate"],
            "unit": right["unit"],
            "features": right_feat["features"],
            "count": right_feat["count"],
            "role": "experiment_ze1",
        },
        "ze1": {
            "filename": right["filename"],
            "signal_name": right["signal_name"],
            "sample_rate": right["sample_rate"],
            "unit": right["unit"],
            "features": right_feat["features"],
            "count": right_feat["count"],
        },
        "pairs": pairs,
        "note": (
            f"ZE1×ZE1 比對（皆 ×{TXT_MV_PER_COUNT} mV/count）；"
            "左欄為對照組 TXT，右欄為實驗組 TXT。"
        ),
    }
    return _attach_correlation_stats(
        result, left=left, right=right, left_feat=left_feat, right_feat=right_feat
    )


def parse_pair_list(
    source: str | Path | BinaryIO,
    *,
    filename: str | None = None,
) -> list[dict[str, Any]]:
    """Parse ZE1 pair-list Excel into row dicts (files not yet resolved)."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise ImportError("需要 openpyxl：pip install openpyxl") from exc

    wb = openpyxl.load_workbook(source, data_only=True)
    ws = wb.active
    rows: list[dict[str, Any]] = []
    for row_i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if row_i == 1:
            continue
        site, ref_name, exp_name, device_note, purpose = (list(row) + [None] * 5)[:5]
        purpose_s = str(purpose or "").strip()
        rows.append(
            {
                "row": row_i,
                "site": str(site or "").strip(),
                "purpose": purpose_s,
                "device_note": str(device_note or "").strip().lower(),
                "ref_listed": str(ref_name).strip() if ref_name else "",
                "exp_listed": str(exp_name).strip() if exp_name else "",
                "expected_count": expected_count_for(purpose_s),
                "source_list": filename or getattr(source, "name", None) or str(source),
            }
        )
    return rows


def resolve_pair_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    index = index_data_files()
    out: list[dict[str, Any]] = []
    for item in rows:
        row = dict(item)
        ref_path = resolve_name(row.get("ref_listed"), index)
        exp_path = resolve_name(row.get("exp_listed"), index)
        row["ref_resolved"] = ref_path.name if ref_path else ""
        row["exp_resolved"] = exp_path.name if exp_path else ""
        row["ref_path"] = str(ref_path) if ref_path else ""
        row["exp_path"] = str(exp_path) if exp_path else ""
        row["files_ok"] = bool(ref_path and exp_path)
        out.append(row)
    return out


def run_one_pair_row(
    row: dict[str, Any],
    *,
    contraction_method: str | None = None,
    feature_method: str | None = None,
) -> dict[str, Any]:
    """Run one resolved pair row; returns summary + full compare result."""
    contr = contraction_method or str(ANALYSIS_CONFIG["contraction_method"])
    feat = feature_method or str(ANALYSIS_CONFIG["feature_method"])
    purpose = str(row.get("purpose") or "")
    device_note = str(row.get("device_note") or "")
    expected = int(row.get("expected_count") or expected_count_for(purpose))

    ref_name = row.get("ref_resolved") or row.get("ref_listed") or ""
    exp_name = row.get("exp_resolved") or row.get("exp_listed") or ""
    if not ref_name or not exp_name:
        return {
            **row,
            "status": "missing_files",
            "note": "缺對照組或實驗組檔案",
            "result": None,
        }

    ref_suffix = Path(ref_name).suffix.lower()
    exp_suffix = Path(exp_name).suffix.lower()
    try:
        if purpose == "刮腿毛" and ref_suffix == ".txt" and exp_suffix == ".txt":
            result = build_feature_compare_txt_pair(
                ref_name,
                exp_name,
                expected_count=expected,
                contraction_method=contr,
                feature_method=feat,
            )
        elif device_note == "ze2" and ref_suffix == ".csv" and exp_suffix == ".txt":
            result = build_feature_compare_ze2(
                ref_name,
                exp_name,
                expected_count=expected,
                contraction_method=contr,
                feature_method=feat,
                ze2_sample_rate=float(ANALYSIS_CONFIG["ze2_sample_rate_hz"]),
                ze2_mv_per_count=float(ANALYSIS_CONFIG["ze2_mv_per_count"]),
                apply_bandpass=True,
            )
        elif ref_suffix == ".csv" and exp_suffix == ".txt":
            result = build_feature_compare(
                ref_name,
                exp_name,
                expected_count=expected,
                contraction_method=contr,
                feature_method=feat,
            )
        elif ref_suffix == ".txt" and exp_suffix == ".txt":
            result = build_feature_compare_txt_pair(
                ref_name,
                exp_name,
                expected_count=expected,
                contraction_method=contr,
                feature_method=feat,
            )
        else:
            raise ValueError(f"不支援的檔案類型組合：{ref_suffix} vs {exp_suffix}")
    except Exception as exc:  # noqa: BLE001
        return {**row, "status": f"error:{exc}", "note": str(exc), "result": None}

    agreement = {
        str(item.get("metric")): item
        for item in (result.get("interval_agreement") or [])
        if isinstance(item, dict) and item.get("metric")
    }
    corr = result.get("correlation") or {}
    rms = agreement.get("rms") or {}
    iemg = agreement.get("iemg") or {}
    return {
        **row,
        "status": "done",
        "contraction_method": contr,
        "feature_method": feat,
        "expected_count": expected,
        "ref_count": (result.get("delsys") or {}).get("count"),
        "exp_count": (
            (result.get("ze2") or result.get("ze1") or result.get("txt") or {}).get("count")
        ),
        "rms_pearson_r": rms.get("pearson_r"),
        "rms_icc": rms.get("icc"),
        "iemg_pearson_r": iemg.get("pearson_r"),
        "iemg_icc": iemg.get("icc"),
        "ttri_rms_r": corr.get("rms"),
        "ttri_iemg_r": corr.get("iemg"),
        "fatigue_visible": (result.get("fatigue_visibility") or {}).get("visible")
        if purpose == "疲勞"
        else "",
        "fatigue_ref": ((result.get("fatigue_visibility") or {}).get("ref") or {}).get("visible")
        if purpose == "疲勞"
        else "",
        "fatigue_exp": ((result.get("fatigue_visibility") or {}).get("exp") or {}).get("visible")
        if purpose == "疲勞"
        else "",
        "note": result.get("note") or "",
        "result": result,
    }


def run_pair_list(
    rows: list[dict[str, Any]],
    *,
    contraction_method: str | None = None,
    feature_method: str | None = None,
    progress: Any | None = None,
) -> list[dict[str, Any]]:
    resolved = resolve_pair_rows(rows)
    out: list[dict[str, Any]] = []
    total = max(len(resolved), 1)
    for idx, row in enumerate(resolved):
        if progress is not None:
            try:
                progress.progress((idx + 1) / total, text=f"分析第 {row.get('row')} 列…")
            except TypeError:
                progress.progress((idx + 1) / total)
        out.append(
            run_one_pair_row(
                row,
                contraction_method=contraction_method,
                feature_method=feature_method,
            )
        )
    return out


def summary_table(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    table: list[dict[str, Any]] = []
    for row in rows:
        table.append(
            {
                "列": row.get("row"),
                "狀態": row.get("status"),
                "部位": row.get("site"),
                "目的": row.get("purpose"),
                "裝置": row.get("device_note"),
                "預期段數": row.get("expected_count"),
                "對照組": row.get("ref_resolved") or row.get("ref_listed"),
                "實驗組": row.get("exp_resolved") or row.get("exp_listed"),
                "段數": f"{row.get('ref_count') or '—'}/{row.get('exp_count') or '—'}",
                "RMS r": row.get("rms_pearson_r"),
                "RMS ICC": row.get("rms_icc"),
                "iEMG r": row.get("iemg_pearson_r"),
                "iEMG ICC": row.get("iemg_icc"),
                "TTRI RMS r": row.get("ttri_rms_r"),
                "疲勞可視": row.get("fatigue_visible") or "—",
                "疲勞(對照)": row.get("fatigue_ref") or "—",
                "疲勞(實驗)": row.get("fatigue_exp") or "—",
            }
        )
    return table


def _safe_sheet_name(name: str, used: set[str]) -> str:
    cleaned = "".join(ch if ch not in r"[]:*?/\"" else "_" for ch in str(name))[:31] or "sheet"
    base = cleaned
    i = 2
    while cleaned in used:
        suffix = f"_{i}"
        cleaned = (base[: 31 - len(suffix)] + suffix)
        i += 1
    used.add(cleaned)
    return cleaned


def _write_sheet_rows(ws: Any, rows: list[dict[str, Any]], fieldnames: list[str] | None = None) -> None:
    if not rows:
        ws.append(["（無資料）"])
        return
    keys = fieldnames or list(rows[0].keys())
    ws.append(keys)
    for row in rows:
        ws.append([row.get(k, "") for k in keys])


def build_pair_list_export_workbook(
    summary_rows: list[dict[str, Any]],
    *,
    result_payloads: list[dict[str, Any]] | None = None,
) -> Any:
    """
    Standalone multi-sheet workbook for list analysis results.

    Sheets:
    - 總覽：summary metrics (+ fatigue columns when present)
    - 疲勞可視：fatigue-only verdict + evidence
    - 區間特徵：per-contraction features for both sides (from JSON payloads)
    """
    try:
        from openpyxl import Workbook
    except ImportError as exc:  # pragma: no cover
        raise ImportError("需要 openpyxl：pip install openpyxl") from exc

    wb = Workbook()
    used: set[str] = set()

    ws_sum = wb.active
    ws_sum.title = _safe_sheet_name("總覽", used)
    overview = summary_table(summary_rows)
    _write_sheet_rows(ws_sum, overview)

    fatigue_rows: list[dict[str, Any]] = []
    for row in summary_rows:
        if str(row.get("purpose") or "") != "疲勞":
            continue
        fatigue_rows.append(
            {
                "列": row.get("row"),
                "狀態": row.get("status"),
                "部位": row.get("site"),
                "裝置": row.get("device_note"),
                "段數": f"{row.get('ref_count') or '—'}/{row.get('exp_count') or '—'}",
                "綜合可視": row.get("fatigue_visible") or "",
                "對照組可視": row.get("fatigue_ref") or "",
                "實驗組可視": row.get("fatigue_exp") or "",
                "對照檔": row.get("ref") or row.get("ref_resolved") or row.get("ref_listed"),
                "實驗檔": row.get("exp") or row.get("exp_resolved") or row.get("exp_listed"),
            }
        )

    # Enrich fatigue sheet with evidence from payloads when available.
    payload_by_row: dict[Any, dict[str, Any]] = {}
    for payload in result_payloads or []:
        if isinstance(payload, dict) and payload.get("row") is not None:
            payload_by_row[payload.get("row")] = payload

    for fr in fatigue_rows:
        payload = payload_by_row.get(fr.get("列")) or {}
        fv = payload.get("fatigue_visibility") or {}
        ref = fv.get("ref") or {}
        exp = fv.get("exp") or {}
        fr["對照依據"] = "; ".join(ref.get("evidence") or [])
        fr["實驗依據"] = "; ".join(exp.get("evidence") or [])
        fr["摘要"] = fv.get("summary") or ""

    ws_fat = wb.create_sheet(_safe_sheet_name("疲勞可視", used))
    _write_sheet_rows(ws_fat, fatigue_rows)

    feature_rows: list[dict[str, Any]] = []
    for payload in result_payloads or []:
        if not isinstance(payload, dict):
            continue
        row_i = payload.get("row")
        site = payload.get("site")
        purpose = payload.get("purpose")
        device = payload.get("device_note")
        pairs = payload.get("pairs") or []
        for pair in pairs:
            if not isinstance(pair, dict):
                continue
            left = pair.get("delsys") or pair.get("ref") or {}
            right = pair.get("txt") or pair.get("ze1") or pair.get("ze2") or pair.get("exp") or {}
            idx = pair.get("index") or left.get("index") or right.get("index")
            feature_rows.append(
                {
                    "列": row_i,
                    "部位": site,
                    "目的": purpose,
                    "裝置": device,
                    "收縮序": idx,
                    "對照_duration": left.get("duration"),
                    "對照_aemg": left.get("aemg"),
                    "對照_rms": left.get("rms"),
                    "對照_iemg": left.get("iemg"),
                    "對照_mpf": left.get("mpf"),
                    "對照_mdf": left.get("mdf"),
                    "對照_peak_rms": left.get("peak_rms"),
                    "實驗_duration": right.get("duration"),
                    "實驗_aemg": right.get("aemg"),
                    "實驗_rms": right.get("rms"),
                    "實驗_iemg": right.get("iemg"),
                    "實驗_mpf": right.get("mpf"),
                    "實驗_mdf": right.get("mdf"),
                    "實驗_peak_rms": right.get("peak_rms"),
                }
            )

    ws_feat = wb.create_sheet(_safe_sheet_name("區間特徵", used))
    _write_sheet_rows(ws_feat, feature_rows)

    agree_rows: list[dict[str, Any]] = []
    for payload in result_payloads or []:
        if not isinstance(payload, dict):
            continue
        for item in payload.get("interval_agreement") or []:
            if not isinstance(item, dict):
                continue
            agree_rows.append(
                {
                    "列": payload.get("row"),
                    "部位": payload.get("site"),
                    "目的": payload.get("purpose"),
                    "裝置": payload.get("device_note"),
                    "指標": item.get("metric"),
                    "n": item.get("n"),
                    "Pearson_r": item.get("pearson_r"),
                    "ICC": item.get("icc"),
                }
            )
    ws_agr = wb.create_sheet(_safe_sheet_name("區間一致性", used))
    _write_sheet_rows(ws_agr, agree_rows)

    return wb


def export_pair_list_results_xlsx(
    summary_rows: list[dict[str, Any]],
    *,
    result_payloads: list[dict[str, Any]] | None = None,
    out_path: Path | None = None,
) -> bytes:
    """Return xlsx bytes; optionally also write to ``out_path``."""
    wb = build_pair_list_export_workbook(summary_rows, result_payloads=result_payloads)
    from io import BytesIO

    buf = BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    if out_path is not None:
        out_path = Path(out_path)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        out_path.write_bytes(data)
    return data


def export_pair_list_results_csv_zip(
    summary_rows: list[dict[str, Any]],
    *,
    result_payloads: list[dict[str, Any]] | None = None,
) -> bytes:
    """Standalone ZIP of CSV sheets mirroring the Excel export."""
    import csv
    import io
    import zipfile

    wb = build_pair_list_export_workbook(summary_rows, result_payloads=result_payloads)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for ws in wb.worksheets:
            text = io.StringIO()
            writer = csv.writer(text)
            for row in ws.iter_rows(values_only=True):
                writer.writerow(["" if v is None else v for v in row])
            zf.writestr(f"{ws.title}.csv", text.getvalue().encode("utf-8-sig"))
    return buf.getvalue()


def load_payloads_from_result_dir(out_dir: Path) -> tuple[list[dict[str, Any]], list[dict[str, Any]]]:
    """Load summary.csv + JSON payloads from an existing pairing_results folder."""
    import csv
    import json

    out_dir = Path(out_dir)
    summary_path = out_dir / "summary.csv"
    summary_rows: list[dict[str, Any]] = []
    if summary_path.exists():
        with summary_path.open(encoding="utf-8", newline="") as fh:
            summary_rows = list(csv.DictReader(fh))

    payloads: list[dict[str, Any]] = []
    for row in summary_rows:
        rel = row.get("json") or ""
        if not rel:
            continue
        path = out_dir / rel
        if path.exists():
            payloads.append(json.loads(path.read_text(encoding="utf-8")))
    return summary_rows, payloads
